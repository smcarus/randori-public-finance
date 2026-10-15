"""One-time extraction of M-13 v0.4 from the Excel workbook into config and data files.

Writes rulebook/*.yaml, data/*.csv and tests/golden/*.json. The golden files hold the
workbook's own calculated values, which the parity tests compare against.

Usage: python scripts/extract_workbook.py path/to/Randori_Country_Risk_Scores_Oct2026.xlsx
"""
import csv
import datetime as dt
import json
import sys
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
R0, LAST = 6, 88
SETTING_KEYS = ['blend', 'coremin', 'popthr', 'pmlthr', 'smallw', 'cap3', 'cap4', 'cap7', 'adjlim', 'gfnadd',
                'pbshock', 'gshock', 'passthru', 'dep', 'resfloat', 'boomboth', 'boomprice', 'squeeze', 'b2bfloor']


def clean(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime('%Y-%m-%d')
    return v


def write_csv(path, header, rows):
    with open(path, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(header)
        for r in rows:
            w.writerow(['' if x is None else clean(x) for x in r])


def main(src):
    wf = openpyxl.load_workbook(src)                 # formulas
    wv = openpyxl.load_workbook(src, data_only=True)  # cached values
    is_const = lambda c: c.value is not None and not (isinstance(c.value, str) and c.value.startswith('='))

    # ---- rulebook ----
    rb = wf['Rulebook']
    inds = []
    for r in range(R0, LAST + 1):
        g = lambda col: rb[f'{col}{r}'].value
        d = {'pillar': g('A'), 'id': g('B'), 'label': g('C'), 'weight': g('D'), 'type': g('E'),
             'thresholds': [g(c) for c in 'FGHIJ']}
        if any(g(c) is not None for c in 'LMNOP'):
            d['high_thresholds'] = [g(c) for c in 'LMNOP']
        inds.append(d)
    ms = wf['Model Settings']
    settings = {k: ms.cell(4 + i, 2).value for i, k in enumerate(SETTING_KEYS)}
    settings['pillar_weights'] = {p: ms.cell(24 + p, 2).value for p in range(1, 8)}
    pillar_names = {p: ms.cell(24 + p, 1).value.split('. ', 1)[1] for p in range(1, 8)}
    ds = wf['Debt Scenarios']
    settings['debt_scenarios'] = {'pb_shock': ds['B5'].value, 'growth_shock': ds['B6'].value,
                                  'depreciation': ds['B7'].value, 'nominal_growth': ds['B8'].value}
    bt = wf['Backtest']
    settings['backtest'] = {'pd_a': bt['D5'].value, 'pd_b': bt['D6'].value, 'include_synthetic': bt['D7'].value,
                            'predictor': bt['D8'].value, 'alpha': bt['D9'].value,
                            'bin_edges': [bt.cell(r, 2).value for r in range(26, 33)]}
    vb = wf['VaR Backtest']
    settings['var_backtest'] = {'confidence': vb['D4'].value, 'include_synthetic': vb['D5'].value, 'alpha': vb['D6'].value}
    mc = wf['Monte Carlo']
    settings['monte_carlo'] = {'score_source': mc['D5'].value, 'rho': mc['D6'].value, 'df': mc['D7'].value,
                               'confidence_levels': [mc['D8'].value, mc['D9'].value, mc['D10'].value],
                               'pd_cap': mc['D13'].value, 'trials': 2000, 'histogram_bands': 20}
    yaml.safe_dump({'version': '0.4', 'pillars': pillar_names, 'indicators': inds},
                   open(ROOT / 'rulebook/rulebook.yaml', 'w'), sort_keys=False, allow_unicode=True, width=120)
    yaml.safe_dump(settings, open(ROOT / 'rulebook/settings.yaml', 'w'), sort_keys=False, width=120)

    # ---- countries and debt scenarios ----
    sc = wv['Scores']
    countries = [sc.cell(r, 2).value for r in range(5, 21)]
    write_csv(ROOT / 'data/countries.csv', ['country', 'region', 'preliminary_score'],
              [[sc.cell(r, 2).value, sc.cell(r, 3).value, sc.cell(r, 4).value] for r in range(5, 21)])
    years = [int(ds.cell(11, c).value) for c in range(4, 21)]
    rows, row_country, cur = [], {}, None
    for r in range(12, ds.max_row + 1):
        if ds.cell(r, 1).value:
            cur = ds.cell(r, 1).value
        if ds.cell(r, 2).value == 'IMF baseline':
            row_country[r] = cur
            fx = ds.cell(r + 2, 3).value
            rows.append([cur, fx] + [ds.cell(r, c).value for c in range(4, 21)])
    write_csv(ROOT / 'data/debt_baseline.csv', ['country', 'fx_share'] + [str(y) for y in years], rows)

    # ---- inputs (constants only; formulas are recomputed by the engine) ----
    mi = wf['Model Inputs']
    ids = {r: mi.cell(r, 2).value for r in range(R0, mi.max_row + 1) if mi.cell(r, 2).value}
    rows = []
    for k, cn in enumerate(countries):
        col = 4 + 3 * k
        assert mi.cell(4, col).value == cn
        for r, iid in ids.items():
            c = mi.cell(r, col)
            if not is_const(c):
                continue
            v = c.value
            if iid == 'r_dsrow':
                iid, v = 'debt_country', row_country[v]
            rows.append([cn, iid, v, mi.cell(r, col + 1).value, mi.cell(r, col + 2).value])
    write_csv(ROOT / 'data/inputs.csv', ['country', 'id', 'value', 'verified', 'source'], rows)

    # ---- scenario drivers ----
    sd = wf['Scenario Drivers']
    drv, ovr = [], []
    for r in range(6, sd.max_row + 1):
        did = sd.cell(r, 1).value
        if not did or ' ' in did:
            continue
        unit = sd.cell(r, 3).value or ''
        kind = 'set' if 'set to' in unit or did == 'd_shock' else ('pct' if did == 'd_res' else 'add')
        label = sd.cell(r, 2).value
        if isinstance(label, str) and label.startswith('='):
            label = wv['Scenario Drivers'].cell(r, 2).value
        drv.append([did, label, unit, kind, sd.cell(r, 4).value, sd.cell(r, 5).value])
        for k, cn in enumerate(countries):
            for s, nm in enumerate(('best', 'worst')):
                v = sd.cell(r, 6 + 2 * k + s).value
                if v is not None:
                    ovr.append([cn, did, nm, v])
    write_csv(ROOT / 'data/scenario_drivers.csv', ['id', 'label', 'unit', 'kind', 'best_default', 'worst_default'], drv)
    write_csv(ROOT / 'data/scenario_overrides.csv', ['country', 'id', 'scenario', 'value'], ovr)

    # ---- backtest logs ----
    rows = []
    for r in range(45, 445):
        v = [bt.cell(r, c).value for c in (1, 2, 3, 4, 5, 6, 8, 10, 11)]
        if v[0] is not None:
            rows.append(v)
    write_csv(ROOT / 'data/backtest_log.csv', ['row_type', 'predictor', 'country', 'as_of', 'horizon_months', 'score',
                                                'outcome', 'note', 'source'], rows)
    rows = []
    for r in range(22, 622):
        v = [vb.cell(r, c).value for c in (1, 2, 3, 4, 5, 8)]
        if v[0] is not None:
            rows.append(v)
    write_csv(ROOT / 'data/var_log.csv', ['row_type', 'portfolio', 'period_end', 'var_forecast', 'realized_loss', 'note'], rows)

    # ---- Monte Carlo ----
    write_csv(ROOT / 'data/mc_portfolio.csv', ['country', 'exposure', 'lgd_mean', 'lgd_sd'],
              [[mc.cell(17 + i, 1).value, mc.cell(17 + i, 5).value, mc.cell(17 + i, 6).value, mc.cell(17 + i, 7).value]
               for i in range(16)])
    dr = wf['MC Draws']
    write_csv(ROOT / 'data/mc_draws.csv', [dr.cell(3, c).value for c in range(1, 2 + 2 + 32)],
              [[dr.cell(r, c).value for c in range(1, 36)] for r in range(4, 2004)])

    # ---- golden values ----
    def num(v):
        return v if isinstance(v, (int, float)) else None
    gold = {}
    for sheet, key in (('Model Calc', 'base'), ('Best Calc', 'best'), ('Worst Calc', 'worst')):
        ws = wv[sheet]
        gold[f'calc_{key}'] = {cn: {ws.cell(r, 2).value: [num(ws.cell(r, 4 + 6 * k).value), num(ws.cell(r, 6 + 6 * k).value)]
                                    for r in range(R0, LAST + 1)} for k, cn in enumerate(countries)}
    hdr = ['wtd_avg', 'worst_core', 'raw', 'triggers', 'adjust', 'floor', 'final']
    for sheet, key in (('Model Pillars', 'base'), ('Best Pillars', 'best'), ('Worst Pillars', 'worst')):
        ws = wv[sheet]
        g = {}
        for i, cn in enumerate(countries):
            r = 6 + i
            g[cn] = {'pillars': {p: {h: num(ws.cell(r, 2 + 7 * (p - 1) + j).value) for j, h in enumerate(hdr)} for p in range(1, 8)},
                     'small_state': ws.cell(r, 51).value, 'score': num(ws.cell(r, 52).value), 'grade': num(ws.cell(r, 53).value)}
            if key == 'base':
                g[cn]['model_on_preliminary'] = num(ws.cell(r, 54).value)
        gold[f'pillars_{key}'] = g
    se = wv['Sensitivity']
    gold['sensitivity'] = {cn: {se.cell(r, 2).value: num(se.cell(r, 5 + 3 * k).value) for r in range(R0, LAST + 1)}
                           for k, cn in enumerate(countries)}
    mcv = wv['Monte Carlo']
    gold['monte_carlo'] = {
        'pd': {mcv.cell(17 + i, 1).value: mcv.cell(17 + i, 4).value for i in range(16)},
        'var': [mcv.cell(r, 3).value for r in (44, 45, 46)], 'es': [mcv.cell(r, 4).value for r in (44, 45, 46)],
        'expected_loss': mcv['C48'].value, 'analytic_expected_loss': mcv['C49'].value, 'sd': mcv['C50'].value,
        'p_one_default': mcv['C51'].value, 'p_two_defaults': mcv['C52'].value, 'max_loss': mcv['C53'].value,
        'es_share': {mcv.cell(17 + i, 1).value: mcv.cell(17 + i, 15).value for i in range(16)},
        'histogram': [mcv.cell(r, 4).value for r in range(59, 79)],
        'default_counts': [mcv.cell(r, 2).value for r in range(82, 89)],
    }
    vbv = wv['VaR Backtest']
    gold['var_backtest_default'] = [vbv.cell(r, 4).value for r in range(8, 16)]
    btv = wv['Backtest']
    gold['backtest_default'] = [btv.cell(r, 4).value for r in range(12, 22)]
    json.dump(gold, open(ROOT / 'tests/golden/workbook_v0.4.json', 'w'), indent=1, default=str)

    # ---- decision modules (v0.5): inputs, settings and golden values ----
    extract_decisions(wf, wv, countries, settings)
    yaml.safe_dump(settings, open(ROOT / 'rulebook/settings.yaml', 'w'), sort_keys=False, width=120)
    print('extracted', len(inds), 'indicators,', len(countries), 'countries')


PERILS = ['ci', 'expropriation', 'political_violence', 'breach', 'nhsfo']


def extract_decisions(wf, wv, countries, settings):
    if 'PRI Perils' not in wf.sheetnames:
        return
    num = lambda v: v if isinstance(v, (int, float)) else None
    dp, pp, rs = wf['Default PD'], wf['PRI Perils'], wf['Restructuring']
    # settings (sections not shown in the workbook keep these defaults)
    settings['default_fit'] = {'features': ['debt_gdp', 'growth', 'inflation_log', 'ca_gdp', 'reserves_months', 'past_default_10y'],
                               'horizon_years': 1, 'bootstrap_draws': 500, 'band': [0.05, 0.95]}
    settings['pri'] = {'perils': {k: {'name': pp.cell(14 + i, 2).value} for i, k in enumerate(PERILS)}}
    settings['pri']['perils']['nhsfo']['drivers'] = 'default_pd'
    settings['restructuring'] = {}
    settings['default_fit'].update(blend_weight_fitted=dp['D5'].value, cds_recovery=dp['D6'].value, gap_tolerance_logodds=dp['D7'].value)
    pri = settings['pri']
    pri.update(mean_reversion=pp['D4'].value, hazard_cap=pp['D5'].value, cost_of_capital=pp['D6'].value, ul_multiplier=pp['D7'].value,
               expense_ratio=pp['D8'].value, discount_rate=pp['D9'].value, score_source=pp['D10'].value)
    for i, k in enumerate(PERILS):
        r = 14 + i; d = pri['perils'][k]
        if k != 'nhsfo':
            d['drivers'] = {p: pp.cell(r, 2 + p).value for p in range(1, 8) if pp.cell(r, 2 + p).value}
            d['base_rate'], d['beta'] = pp.cell(r, 10).value, pp.cell(r, 11).value
        d['persistence'], d['severity'], d['recovery'] = pp.cell(r, 12).value, pp.cell(r, 13).value, pp.cell(r, 14).value
        d['sector_multiplier'] = pp.cell(r, 15).value == 1
    r = 22
    sm = {}
    while pp.cell(r, 1).value:
        sm[pp.cell(r, 1).value] = pp.cell(r, 4).value; r += 1
    pri['sector_multipliers'] = sm
    r += 2
    of = {}
    while pp.cell(r, 1).value:
        of[pp.cell(r, 1).value] = pp.cell(r, 4).value; r += 1
    pri['obligor_factors'] = of
    rk = ['valuation_year', 'base_months', 'china_share_threshold', 'china_months', 'no_cac_months', 'series_cac_months',
          'multi_law_months', 'imf_program_months', 'common_framework_months']
    settings['restructuring'].update({k: rs.cell(4 + i, 4).value for i, k in enumerate(rk)})
    # inputs
    terms = ['intercept', 'debt_gdp', 'growth', 'inflation_log', 'ca_gdp', 'reserves_months', 'past_default_10y']
    write_csv(ROOT / 'data/default_coefficients.csv', ['term', 'coefficient', 'note'],
              [[t, dp.cell(10 + i, 4).value, 'from the Default PD tab'] for i, t in enumerate(terms)])
    write_csv(ROOT / 'data/default_features_extra.csv', ['country', 'reserves_months', 'past_default_10y', 'source'],
              [[cn, dp.cell(20 + i, 9).value, dp.cell(20 + i, 10).value, ''] for i, cn in enumerate(countries)])
    write_csv(ROOT / 'data/default_bands.csv', ['country', 'pd_low', 'pd_high', 'note'],
              [[cn, dp.cell(20 + i, 12).value, dp.cell(20 + i, 13).value, ''] for i, cn in enumerate(countries)])
    write_csv(ROOT / 'data/market_spreads.csv', ['country', 'cds_5y_bps', 'as_of', 'source'],
              [[cn, dp.cell(20 + i, 15).value, dp.cell(20 + i, 16).value, ''] for i, cn in enumerate(countries)])
    ck = ['multilateral', 'paris_club', 'china', 'other_bilateral', 'bonds', 'commercial_banks', 'domestic', 'imf_program', 'common_framework', 'as_of', 'source']
    write_csv(ROOT / 'data/creditor_map.csv', ['country'] + ck, [[rs.cell(17 + i, 1).value] + [rs.cell(17 + i, j).value for j in range(2, 13)]
                                                                 for i in range(len(countries))])
    ph = ['policy_id', 'country', 'sector', 'obligor', 'term_years', 'limit', 'insured_share', 'deductible', 'location_mult'] + PERILS + ['note']
    po = [r for r in range(1, pp.max_row + 1) if pp.cell(r, 1).value == 'Policy'][0]
    write_csv(ROOT / 'data/pri_policies.csv', ph, [[pp.cell(r, j).value for j in range(1, 16)] for r in range(po + 1, po + 21) if pp.cell(r, 1).value])
    bh = ['bond_id', 'country', 'face', 'coupon', 'maturity', 'amort_start', 'governing_law', 'cac', 'price', 'face_haircut', 'new_coupon',
          'new_maturity', 'grace_years', 'exit_yield', 'note']
    write_csv(ROOT / 'data/bonds.csv', bh, [[rs.cell(r, j).value for j in range(1, 16)] for r in range(37, 57) if rs.cell(r, 1).value])
    # golden
    dv, pv_, rv = wv['Default PD'], wv['PRI Perils'], wv['Restructuring']
    gold = {'default': {cn: {'scorecard_pd': num(dv.cell(20 + i, 4).value), 'fitted_pd': num(dv.cell(20 + i, 11).value),
                             'blended_pd': num(dv.cell(20 + i, 14).value), 'market_pd': num(dv.cell(20 + i, 17).value)}
                        for i, cn in enumerate(countries)},
            'policies': {pv_.cell(r, 1).value: [num(pv_.cell(r, j).value) for j in range(16, 30)] for r in range(po + 1, po + 21) if pv_.cell(r, 1).value},
            'bonds': {rv.cell(r, 1).value: [num(rv.cell(r, j).value) for j in (16, 17, 18, 19, 20, 22, 23)] for r in range(37, 57) if rv.cell(r, 1).value}}
    json.dump(gold, open(ROOT / 'tests/golden/workbook_v0.5_decisions.json', 'w'), indent=1)


if __name__ == '__main__':
    main(sys.argv[1])
