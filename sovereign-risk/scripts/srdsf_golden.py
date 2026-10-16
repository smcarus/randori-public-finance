"""Save every cell of the MAC SRDSF template (inputs and the values Excel cached for formulas) for the parity tests.

    python scripts/srdsf_golden.py path/to/mac-srdsf-template.xlsm [outdir]

Writes tests/golden/srdsf/<sheet>.json.gz with {sheet, cells: [[row, col, value]], formulas: [[row, col]]} and
names.json with the workbook's defined names (lower case -> 'Sheet!$A$1'). The template ships with the IMF's
fictional "MAC DSA sample" country, so these values are the template's own results for that sample.
"""
import datetime
import gzip
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

ROOT = Path(__file__).resolve().parents[1]


def _val(v):
    if isinstance(v, (datetime.date, datetime.datetime, datetime.time, datetime.timedelta)):
        return str(v) if isinstance(v, datetime.timedelta) else v.isoformat()
    return v


def main(src, out=None):
    out = Path(out or ROOT / 'tests' / 'golden' / 'srdsf')
    out.mkdir(parents=True, exist_ok=True)
    wf = openpyxl.load_workbook(src)
    wv = openpyxl.load_workbook(src, data_only=True)
    names = {n.lower(): dn.attr_text for n, dn in wf.defined_names.items()}
    for ws in wf.worksheets:
        for n, dn in ws.defined_names.items():
            names.setdefault(n.lower(), dn.attr_text)
    (out / 'names.json').write_text(json.dumps(names, indent=0, sort_keys=True))
    for ws in wf.worksheets:
        v = wv[ws.title]
        cells, formulas = [], []
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                if isinstance(c.value, ArrayFormula) or (isinstance(c.value, str) and c.value.startswith('=')):
                    formulas.append([c.row, c.column])
                    cells.append([c.row, c.column, _val(v.cell(c.row, c.column).value)])
                else:
                    cells.append([c.row, c.column, _val(c.value)])
        data = json.dumps({'sheet': ws.title, 'cells': cells, 'formulas': formulas}, separators=(',', ':'))
        (out / f'{ws.title}.json.gz').write_bytes(gzip.compress(data.encode(), mtime=0))
        print(ws.title, len(cells), 'cells', len(formulas), 'formulas')


if __name__ == '__main__':
    main(*sys.argv[1:])
