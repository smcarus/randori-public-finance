"""Independent oracle for the LIC DSF runner: evaluates the template's own formulas, read from the .xlsm with openpyxl,
over a Book of input values.

LibreOffice cannot recalculate the LIC DSF template, so changed-input checks need another evaluator. This one runs
every formula cell of the workbook through the general Excel evaluator of sovrisk.licdsf.inputs (Engine, written
for the label / look-up sheets, not from the hand-ported sheet code), extended here with the functions it lacks
(CHOOSE, NA, NORMDIST, RANK, HLOOKUP, COUNTIF, LARGE, STDEV, MIN, SEARCH, INDIRECT, ISTEXT). Evaluation is lazy and
cell by cell, so it needs no sheet order. Cells sharing a formula pattern (same formula in relative R1C1 form) are
compiled once.

    formulas = read_formulas(XLSM)              # {sheet: {pattern key: (kind, [(row, col), ...], text)}}
    evaluate_all(book, formulas)                # fills every formula cell of book from its input cells
"""
import math
import re
import statistics
import sys
import threading
from pathlib import Path

from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.datetime import from_excel

from sovrisk.licdsf import inputs as I
from sovrisk.srdsf.book import ci, cl

XLSM = Path('/mnt/project-files/inputs/licdsf/LIC-DSF-IDA21-Template-08-12-2025-vf.xlsm')

_REF = re.compile(r"(?<![A-Za-z0-9_.])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![A-Za-z0-9_(])")


def _key(text, r, c):
    """The formula in relative R1C1-like form (string literals left alone)."""
    parts = text.split('"')
    for i in range(0, len(parts), 2):
        def sub(m):
            col = ci(m.group(2))
            rr = int(m.group(4))
            cs = f'C{col}' if m.group(1) else f'C[{col - c}]'
            rs = f'R{rr}' if m.group(3) else f'R[{rr - r}]'
            return rs + cs
        parts[i] = _REF.sub(sub, parts[i])
    return '"'.join(parts)


def read_formulas(path=XLSM):
    """{sheet: {key: (kind, cells, text)}}: every formula cell of the workbook grouped by pattern; kind is 'array'
    for a CSE formula, 'date' for a date-formatted cell, else None; text is the anchor (first) cell's formula."""
    import warnings
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        wb = openpyxl.load_workbook(path)
    out = {}
    for ws in wb.worksheets:
        groups = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, ArrayFormula):
                    kind, text = 'array', v.text
                elif isinstance(v, str) and v.startswith('='):
                    kind, text = ('date' if is_date_format(c.number_format) else None), v
                else:
                    continue
                k = (kind, _key(text, c.row, c.column))
                if k not in groups:
                    groups[k] = (kind, [], text)
                groups[k][1].append((c.row, c.column))
        out[ws.title] = groups
    return out


def _normcdf(x, mu, sd):
    return 0.5 * (1 + math.erf((x - mu) / (sd * math.sqrt(2))))


_PAREN_RANGE = re.compile(r'(\$?[A-Z]{1,3}\$?\d+):\(OFFSET\(\1,(\d+),0\)\)')


class Oracle(I.Engine):
    """inputs.Engine plus the Excel functions the full template uses beyond the sheets inputs.py ports."""

    def compile(self, formula, array=False):
        # Chart Data!AM: RANK(AF47,$AF47:(OFFSET($AF47,3,0))) - the tokenizer reads 'A1:(' as a function name, so
        # the range is rewritten as OFFSET($AF47,0,0,4,1) (same reference)
        formula = _PAREN_RANGE.sub(lambda m: f'OFFSET({m.group(1)},0,0,{int(m.group(2)) + 1},1)', formula)
        return super().compile(formula, array)

    def op_node(self, op, fa, fb):
        """inputs.Engine's operators, with Excel's 15-digit tolerance: a sum or difference that cancels to within
        1e-15 of its operands is 0, and comparisons treat numbers that equal to 15 significant digits as equal.
        An operand that fails (text in arithmetic raises inside inputs.Engine) is an error value, so with two
        failing operands the left one wins, as in Excel (=1/0*("a"+1) is #DIV/0!)."""
        fa, fb = self._caught(fa), self._caught(fb)
        if op in ('+', '-'):
            def g(a, b):
                x, y = I.to_num(a), I.to_num(b)
                z = x + y if op == '+' else x - y
                return 0.0 if z != 0 and abs(z) <= 1e-15 * max(abs(x), abs(y)) else z
            return lambda ctx: I.lift2(g, fa(ctx), fb(ctx))
        if op in ('=', '<>', '<', '>', '<=', '>='):
            def g(a, b):
                if I.is_err(a):
                    raise I.XlError(a)
                if I.is_err(b):
                    raise I.XlError(b)
                num = (int, float)
                if isinstance(a, num) and isinstance(b, num) and not isinstance(a, bool) and not isinstance(b, bool) \
                        and abs(a - b) <= 1e-15 * max(abs(a), abs(b)):
                    c = 0
                else:
                    c = I.compare(a, b)
                return {'=': c == 0, '<>': c != 0, '<': c < 0, '>': c > 0, '<=': c <= 0, '>=': c >= 0}[op]
            return lambda ctx: I.lift2(g, fa(ctx), fb(ctx))
        return super().op_node(op, fa, fb)

    @staticmethod
    def _caught(f):
        def g(ctx):
            try:
                return f(ctx)
            except I.XlError as e:
                return e.e
        return g

    def f_NA(self, args, arr):
        return lambda ctx: I.NA

    def f_CHOOSE(self, args, arr):
        i = self.scalar(args[0], arr)
        opts = [self.node(a, arr) for a in args[1:]]

        def f(ctx):
            k = i(ctx)
            if I.is_err(k):
                return k
            k = int(I.to_num(k))
            if k < 1 or k > len(opts):
                return I.VALUE
            return opts[k - 1](ctx)
        return f

    def f_ISTEXT(self, args, arr):
        return self._is(args, arr, lambda v: isinstance(v, str) and not I.is_err(v))

    def f_MIN(self, args, arr):
        return self._agg(args, lambda xs: min(xs) if xs else 0)

    def f_LARGE(self, args, arr):
        g, k = self.as_array(args[0]), self.scalar(args[1], arr)

        def f(ctx):
            xs = self._nums([g], ctx)
            n = int(I.to_num(k(ctx)))
            if n < 1 or n > len(xs):
                return I.NUM
            return sorted(xs, reverse=True)[n - 1]
        return f

    def f_STDEV(self, args, arr):
        return self._agg(args, lambda xs: statistics.stdev(xs) if len(xs) > 1 else I.DIV0)

    def f_NORMDIST(self, args, arr):
        def nd(x, mu, sd, cum):
            x, mu, sd = I.to_num(x), I.to_num(mu), I.to_num(sd)
            if sd <= 0:
                return I.NUM
            if I.to_bool(cum):
                return _normcdf(x, mu, sd)
            return math.exp(-0.5 * ((x - mu) / sd) ** 2) / (sd * math.sqrt(2 * math.pi))
        return self._scalar_fn(args, arr, nd, 4)

    def f_RANK(self, args, arr):
        x, g = self.scalar(args[0], arr), self.as_array(args[1])
        order = self._opt(args, 2, arr, 0)

        def f(ctx):
            v = x(ctx)
            if I.is_err(v):
                return v
            vals = self.flat(g(ctx))
            nums = [y for y in vals if isinstance(y, (int, float)) and not isinstance(y, bool)]
            v = I.to_num(v)
            if v not in nums:
                return I.NA
            asc = I.to_num(order(ctx) or 0) != 0
            return 1 + sum(1 for y in nums if (y < v if asc else y > v))
        return f

    def f_COUNTIF(self, args, arr):
        g, crit = self.as_array(args[0]), self.scalar(args[1], arr)

        def f(ctx):
            c = crit(ctx)
            if I.is_err(c):
                return c
            m = re.match(r'^(<=|>=|<>|<|>|=)?(.*)$', I.to_text(c), re.S)
            op, rhs = m.group(1) or '=', m.group(2)
            try:
                rv = float(rhs)
            except ValueError:
                rv = rhs
            n = 0
            for y in self.flat(g(ctx)):
                if isinstance(rv, float):
                    if not isinstance(y, (int, float)) or isinstance(y, bool):
                        if op == '<>':
                            n += 1
                        continue
                    d = (y > rv) - (y < rv)
                else:
                    if y is None or not isinstance(y, str):
                        if op == '<>':
                            n += 1
                        continue
                    a, b = y.lower(), rv.lower()
                    d = (a > b) - (a < b)
                n += {'=': d == 0, '<>': d != 0, '<': d < 0, '>': d > 0, '<=': d <= 0, '>=': d >= 0}[op]
            return n
        return f

    def f_HLOOKUP(self, args, arr):
        x, g, row = self.scalar(args[0], arr), self.as_array(args[1]), self.scalar(args[2], arr)
        approx = self._opt(args, 3, arr, True)

        def f(ctx):
            v = x(ctx)
            if I.is_err(v):
                return v
            t = g(ctx)
            if not isinstance(t, I.Ref):
                return I.VALUE
            first = [self.cell(t.sheet, t.r0, c) for c in range(t.c0, t.c1 + 1)]
            i = I.xl_match(v, first, 1 if I.to_bool(approx(ctx)) else 0)
            if I.is_err(i):
                return i
            k = int(I.to_num(row(ctx)))
            if k < 1 or k > t.h:
                return I.REF
            return self.cell(t.sheet, t.r0 + k - 1, t.c0 + i - 1)
        return f

    def f_SEARCH(self, args, arr):
        def search(find, within, start=1):
            i = I.to_text(within).lower().find(I.to_text(find).lower(), int(I.to_num(start or 1)) - 1)
            return I.VALUE if i < 0 else i + 1
        return self._scalar_fn(args, arr, search, 3)

    def f_CELL(self, args, arr):
        if len(args) < 2:                   # CELL("filename") of the active sheet: #VALUE! in the cached template
            return lambda ctx: I.VALUE
        return super().f_CELL(args, arr)

    def f_INDIRECT(self, args, arr):
        g = self.scalar(args[0], arr)

        def f(ctx):
            v = g(ctx)
            if I.is_err(v):
                return v
            try:
                return self.compile('=' + I.to_text(v))(ctx)
            except Exception:
                return I.REF
        return f


def evaluate_all(book, formulas, sheets=None):
    """Blank every formula cell of the book (of `sheets`, default all) and compute it from the template's own
    formulas. Returns the Engine."""
    eng = Oracle(book)
    dates = []
    for sheet, groups in formulas.items():
        if sheets is not None and sheet not in sheets:
            continue
        S = book[sheet]
        for kind, cells, text in groups.values():
            for r, c in cells:
                S.cells.pop((r, c), None)
            try:
                eng.add(sheet, ' '.join(f'{cl(c)}{r}' for r, c in cells), text, kind if kind == 'array' else None)
                if kind == 'date':
                    dates += [(sheet, r, c) for r, c in cells]
            except SyntaxError:             # text typed with a leading '=' (PV_Base-add.cost.mkt!A3): a constant
                for r, c in cells:
                    S.set(r, c, text)
    box = {}

    def go():
        try:
            eng.run()
        except BaseException as e:          # re-raised in the caller's thread
            box['err'] = e
    sys.setrecursionlimit(1_000_000)
    old = threading.stack_size(1 << 30)
    try:
        t = threading.Thread(target=go)
        t.start()
        t.join()
    finally:
        threading.stack_size(old)
    if 'err' in box:
        raise box['err']
    for sheet, r, c in dates:               # stored as openpyxl reads a date-formatted cell (golden convention)
        v = book[sheet].get(r, c)
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            book[sheet].set(r, c, from_excel(v).isoformat())
    return eng
